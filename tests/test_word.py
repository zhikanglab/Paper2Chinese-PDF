"""Real DOCX output, native math runs and dual-format smoke regression."""
import json
from argparse import Namespace
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from toolkit.cli import audit, build


def test_both_formats_and_true_word_subscripts(tmp_path: Path):
    source = tmp_path / 'source.pdf'
    src = fitz.open()
    page = src.new_page(width=595, height=842)
    page.insert_text((100, 130), 'Plot 1', fontsize=15)
    page.draw_rect(fitz.Rect(95, 145, 495, 295))
    src.save(str(source)); src.close()
    manifest = tmp_path / 'paper.json'
    manifest.write_text(json.dumps({
        'source_pdf': 'source.pdf',
        'metadata': {'title_zh': '合成学术中文译文', 'title_en': 'Synthetic scientific source'},
        'blocks': [
            {'type': 'heading', 'text': '摘要'},
            {'type': 'paragraph', 'source_page': 1, 'source_paragraph_id': 'p1',
             'runs': [{'text': '能量项'}, {'math': [{'text': 'W'}, {'sub': 'I'}]},
                      {'text': '与'}, {'math': [{'text': 'CO'}, {'sub': '2'}]},
                      {'text': '及'}, {'math': [{'text': 's'}, {'sup': '−1'}]},
                      {'text': '均依原文保持上下标。'}]},
            {'type': 'figure', 'source_page': 1, 'bbox': [95, 145, 495, 295], 'label': '图1', 'caption': '合成示意图'},
        ],
        'expected_counts': {'figures': 1, 'tables': 0, 'equations': 0, 'references': 0},
        'quality': {'translation_reviewed': False, 'layout_reviewed': False},
    }, ensure_ascii=False), encoding='utf-8')
    assert build(Namespace(manifest=str(manifest), out=str(tmp_path / 'translated'), format='both')) == 0
    pdf = tmp_path / 'translated.pdf'
    docx = tmp_path / 'translated.docx'
    assert pdf.is_file() and docx.is_file()
    doc = Document(docx)
    assert doc.paragraphs[0].text == '合成学术中文译文'
    assert doc.paragraphs[1].text == 'Synthetic scientific source'
    assert len(doc.inline_shapes) == 1
    subs = [r.text for p in doc.paragraphs for r in p.runs if r.font.subscript]
    sups = [r.text for p in doc.paragraphs for r in p.runs if r.font.superscript]
    assert 'I' in subs and '2' in subs and '−1' in sups
    assert all(p.alignment not in (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.DISTRIBUTE) for p in doc.paragraphs)
    assert audit(Namespace(manifest=str(manifest), pdf=str(docx), strict=False)) == 0
    assert audit(Namespace(manifest=str(manifest), pdf=str(pdf), strict=False)) == 0


def test_plain_underscore_science_not_accepted(tmp_path: Path):
    source = tmp_path / 'source.pdf'
    src = fitz.open(); src.new_page(); src.save(source); src.close()
    manifest = tmp_path / 'p.json'
    manifest.write_text(json.dumps({
        'source_pdf': 'source.pdf',
        'metadata': {'title_zh': '测试', 'title_en': 'Test'},
        'blocks': [{'type':'paragraph', 'source_page':1, 'text':'透射能量 W_I 的结果'}],
        'quality': {},
    }, ensure_ascii=False), encoding='utf-8')
    assert build(Namespace(manifest=str(manifest), out=str(tmp_path / 'fail.docx'), format='docx')) == 2
    assert not (tmp_path / 'fail.docx').exists()
