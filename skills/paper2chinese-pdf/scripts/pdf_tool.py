"""Inspectable PDF layout helpers. This module does NOT translate language."""
from __future__ import annotations

import argparse
import csv
import re
import html
import json
import sys
from collections import Counter
from pathlib import Path

import fitz
from weasyprint import HTML


CSS = """
@page { size: A4; margin: 19mm 20mm 17mm 20mm; }
html { color: #15171a; }
body { font-family: 'Noto Sans CJK SC', 'Source Han Sans SC', 'Microsoft YaHei', sans-serif;
       font-size: 10.6pt; line-height: 1.53; font-kerning: normal; }
* { letter-spacing: normal; word-spacing: normal; }
h1 { font-size: 17.2pt; line-height: 1.35; font-weight: 700; margin: 0 0 2.8mm 0; }
.original-title { font-family: sans-serif; color: #4a4c51; font-size: 10.8pt;
                   font-weight: 400; margin-bottom: 4mm; line-height: 1.4; }
.meta { font-size: 9.3pt; color: #666; margin-bottom: 8mm; }
h2 { font-size: 12.7pt; margin-top: 6mm; margin-bottom: 2mm; break-after: avoid; }
h3 { font-size: 11.5pt; margin-top: 4mm; margin-bottom: 2mm; break-after: avoid; }
/* No full justification: CJK+Latin prose must never be stretched to fill lines. */
p { margin: 0 0 3.1mm 0; text-align: left; text-align-last: left;
    line-break: strict; word-break: normal; overflow-wrap: anywhere;
    orphans: 2; widows: 2; }
.math-token { font-family: 'STIX Two Math', 'Noto Sans Math', 'DejaVu Sans', serif;
    font-style: normal; white-space: nowrap; }
.math-token sub, .math-token sup { font-size: .72em; line-height: 0;
    position: relative; vertical-align: baseline; }
.math-token sub { bottom: -.32em; }
.math-token sup { top: -.48em; }

.fig, .tbl, .eq { margin: 3mm auto 4.5mm auto; text-align: center;
      break-inside: avoid; page-break-inside: avoid; }
.fig img, .tbl img, .eq img { max-width: 100%; max-height: 223mm;
      width: auto; height: auto; object-fit: contain; }
.fig figcaption, .tbl figcaption { margin: 1.2mm auto 0; font-size: 9.4pt;
                  color: #333; line-height: 1.35; text-align: center; }
.eq figcaption { margin-top: 0.7mm; font-size: 9.1pt; text-align: right; }
.ref { font-size: 8.9pt; color: #282828; line-height: 1.35; margin-bottom: 1.6mm; text-align: left; }
.narrative { margin: 2.5mm auto; }
.callout { font-size: 9.4pt; border-left: 2px solid #bbb; padding-left: 3mm;
           color: #484848; margin: 3mm 0; }
"""


def _read_manifest(path: Path) -> dict:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        raise ValueError(f"Cannot read manuscript JSON: {e}") from e
    if not isinstance(payload, dict):
        raise ValueError("Manifest root must be a JSON object")
    return payload


def _source_path(manifest: dict, manifest_file: Path) -> Path:
    name = manifest.get("source_pdf")
    if not name:
        raise ValueError("Missing source_pdf in manifest")
    p = Path(name)
    if not p.is_absolute():
        p = manifest_file.parent / p
    if not p.is_file():
        raise ValueError(f"Source PDF not found: {p}")
    return p.resolve()


def _e(v: object) -> str:
    return html.escape(str(v), quote=True)


# High-risk plaintext expressions: use semantic math runs (or original formula crop).
_BAD_MATH = re.compile(r"(?<![A-Za-z])(?:W_[A-Za-z0-9]+|CO_2|H_2O|NaHCO_3|kg/m\^3|m\^2|m\^3|s\^-1|s\^\{-1\})")


def _math_html(parts: list[dict]) -> str:
    if not isinstance(parts, list) or not parts:
        raise ValueError("Math must contain nonempty run parts")
    chunks = []
    for part in parts:
        if not isinstance(part, dict) or len(part) != 1:
            raise ValueError("Math parts must each contain text, sub or sup")
        key, val = next(iter(part.items()))
        if key not in {"text", "sub", "sup"} or not str(val):
            raise ValueError("Invalid math run (only text/sub/sup supported)")
        safe = _e(val)
        chunks.append(f"<{key}>{safe}</{key}>" if key != "text" else safe)
    return '<span class="math-token">' + ''.join(chunks) + '</span>'


def _runs_html(block: dict, field: str = "runs", fallback: str = "text") -> str:
    runs = block.get(field)
    if runs is None:
        return _e(block.get(fallback, ""))
    if not isinstance(runs, list) or not runs:
        raise ValueError(f"{field} must be nonempty list")
    out = []
    for run in runs:
        if not isinstance(run, dict) or len(run) != 1:
            raise ValueError("Each run must contain one key (text or math)")
        key, val = next(iter(run.items()))
        if key == "text":
            out.append(_e(val))
        elif key == "math":
            out.append(_math_html(val))
        else:
            raise ValueError(f"Unsupported run kind: {key}")
    return ''.join(out)


def _plain_runs(block: dict) -> str:
    if "runs" not in block:
        return str(block.get("text", ""))
    chunks = []
    for run in block["runs"]:
        if "text" in run: chunks.append(str(run["text"]))
        elif "math" in run:
            for piece in run["math"]:
                k,v = next(iter(piece.items()))
                chunks.append({"text":"", "sub":"₍", "sup":"⁽"}[k] + str(v) + ("₎" if k=="sub" else ("⁾" if k=="sup" else "")))
    return ''.join(chunks)


def _crop(doc: fitz.Document, block: dict, img_dir: Path, serial: int) -> Path:
    number = block.get("source_page")
    bbox = block.get("bbox")
    if not isinstance(number, int) or not 1 <= number <= len(doc):
        raise ValueError(f"Block {serial}: source_page must be 1..{len(doc)}")
    if not isinstance(bbox, list) or len(bbox) != 4:
        raise ValueError(f"Block {serial}: bbox must contain 4 PDF point coordinates")
    rect = fitz.Rect(*[float(x) for x in bbox])
    page = doc[number - 1]
    if rect.is_empty or not page.rect.contains(rect):
        raise ValueError(f"Block {serial}: bbox {bbox} outside source page {number}")
    img_dir.mkdir(parents=True, exist_ok=True)
    output = img_dir / f"asset_{serial:03d}.png"
    pix = page.get_pixmap(clip=rect, matrix=fitz.Matrix(3, 3), alpha=False)
    pix.save(str(output))
    return output


def inspect_pdf(args: argparse.Namespace) -> int:
    p = Path(args.pdf)
    if not p.is_file():
        print(f"PDF not found: {p}", file=sys.stderr)
        return 2
    with fitz.open(str(p)) as doc:
        result = {
            "source_pdf": str(p.resolve()),
            "pages": len(doc),
            "notes": ["Reading order in two-column PDFs is not guaranteed; examine page previews and original PDF."],
            "page_inventory": [],
        }
        previews = Path(args.preview) if args.preview else None
        if previews:
            previews.mkdir(parents=True, exist_ok=True)
        for i, page in enumerate(doc):
            blocks = page.get_text("blocks", sort=True)
            entry = {"page": i + 1, "width": round(page.rect.width, 1),
                     "height": round(page.rect.height, 1),
                     "text_blocks": [{"bbox": [round(float(x), 1) for x in b[:4]],
                                      "preview": str(b[4][:280]).strip()}
                                     for b in blocks if len(b) > 4 and str(b[4]).strip()],
                     "embedded_images": len(page.get_images())}
            result["page_inventory"].append(entry)
            if previews:
                pix = page.get_pixmap(matrix=fitz.Matrix(0.90, 0.90), alpha=False)
                pix.save(str(previews / f"page_{i + 1:03d}.png"))
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Inspected {result['pages']} pages: {out}")
    return 0


def _check_manifest(d: dict, source: fitz.Document) -> tuple[list[str], list[str]]:
    errs, warns = [], []
    meta = d.get("metadata", {})
    if not str(meta.get("title_zh", "")).strip():
        errs.append("Missing metadata.title_zh")
    if not str(meta.get("title_en", "")).strip():
        errs.append("Missing exact metadata.title_en")
    blocks = d.get("blocks", [])
    if not isinstance(blocks, list) or not blocks:
        errs.append("Missing ordered content blocks")
        return errs, warns
    allowed = {"heading", "paragraph", "figure", "table", "equation", "references", "note"}
    for i, b in enumerate(blocks, start=1):
        kind = b.get("type")
        if kind not in allowed:
            errs.append(f"Block {i}: unknown type {kind}")
            continue
        if kind in {"heading", "paragraph", "note"}:
            try:
                content = _plain_runs(b) if kind == "paragraph" else str(b.get("text", ""))
                if not content.strip(): errs.append(f"Block {i}: empty {kind}")
                if kind == "paragraph" and "runs" not in b and _BAD_MATH.search(content):
                    errs.append(f"Block {i}: flattened scientific sub/sup notation; use semantic math runs")
                if kind == "paragraph": _runs_html(b)
            except (ValueError, TypeError) as exc:
                errs.append(f"Block {i}: {exc}")
        if kind == "paragraph" and "source_page" in b:
            if not isinstance(b["source_page"], int) or not 1 <= b["source_page"] <= len(source):
                errs.append(f"Block {i}: invalid paragraph source_page")
        if kind == "paragraph" and not b.get("source_page"):
            warns.append(f"Block {i}: missing source_page; source-to-translation traceability is limited")
        if kind == "references" and not b.get("items"):
            errs.append(f"Block {i}: missing reference items")
        if kind in {"figure", "table", "equation"}:
            try:
                page = source[int(b["source_page"]) - 1]
                rect = fitz.Rect(*[float(v) for v in b["bbox"]])
                if rect.is_empty or not page.rect.contains(rect):
                    errs.append(f"Block {i}: crop bbox outside source page")
            except (KeyError, IndexError, TypeError, ValueError):
                errs.append(f"Block {i}: missing/invalid crop coordinates")
            if kind in {"figure", "table"} and not b.get("caption"):
                warns.append(f"Block {i}: missing Chinese caption")
    expected = d.get("expected_counts", {})
    count = Counter(b.get("type") for b in blocks)
    for expected_key, actual_key in (("figures", "figure"), ("tables", "table"), ("equations", "equation")):
        if expected_key in expected and count[actual_key] != expected[expected_key]:
            errs.append(f"Expected {expected[expected_key]} {expected_key}; found {count[actual_key]} {actual_key} blocks")
    if "references" in expected:
        ref_count = sum(len(b.get("items", [])) for b in blocks if b.get("type") == "references")
        if ref_count != expected["references"]:
            errs.append(f"Expected {expected['references']} bibliography entries; found {ref_count}")
    if not any(b.get("type") == "paragraph" for b in blocks):
        errs.append("No translated full-text paragraphs")
    return errs, warns


def _build_pdf(args: argparse.Namespace) -> int:
    mf = Path(args.manifest).resolve()
    d = _read_manifest(mf)
    with fitz.open(str(_source_path(d, mf))) as source:
        errors, warnings = _check_manifest(d, source)
        if errors:
            for m in errors: print("ERROR:", m, file=sys.stderr)
            return 2
        for m in warnings: print("WARNING:", m, file=sys.stderr)
        out = Path(args.out).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        asset_dir = out.parent / (out.stem + "_draft_assets")
        meta = d.get("metadata", {})
        pieces = ['<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><style>',
                  CSS, '</style></head><body>', f'<h1>{_e(meta["title_zh"])}</h1>',
                  f'<div class="original-title">{_e(meta["title_en"])}</div>']
        subtitle = ' · '.join(str(meta.get(x)) for x in ("authors", "journal", "year", "doi") if meta.get(x))
        if subtitle:
            pieces.append(f'<div class="meta">{_e(subtitle)}</div>')
        for i, b in enumerate(d["blocks"], start=1):
            kind = b["type"]
            if kind == "heading":
                level = 2 if int(b.get("level", 2)) == 2 else 3
                pieces.append(f'<h{level}>{_e(b["text"])}</h{level}>')
            elif kind == "paragraph":
                if "runs" in b:
                    pieces.append(f'<p>{_runs_html(b)}</p>')
                else:
                    paras = [p for p in str(b["text"]).split("\n\n") if p.strip()]
                    pieces.extend(f'<p>{_e(p).replace(chr(10), "<br>")}</p>' for p in paras)
            elif kind == "note":
                pieces.append(f'<div class="callout">译者核对注：{_e(b["text"])}</div>')
            elif kind == "references":
                pieces.extend(f'<p class="ref">{_e(x)}</p>' for x in b["items"])
            else:
                target = _crop(source, b, asset_dir, i).resolve().as_uri()
                cls = {"figure": "fig", "table": "tbl", "equation": "eq"}[kind]
                width = float(b["bbox"][2]) - float(b["bbox"][0])
                img_width_mm = max(24, min(164, width * 0.3527))
                pieces.append(f'<figure class="{cls}"><img src="{_e(target)}" style="max-width: {img_width_mm:.1f}mm" alt="{_e(b.get("label", kind))}">')
                caption = b.get("caption", "")
                if kind != "equation" and (caption or b.get("label")):
                    pieces.append(f'<figcaption>{_e(b.get("label", ""))} {_e(caption)}</figcaption>')
                elif kind == "equation" and b.get("show_label", False):
                    pieces.append(f'<figcaption>{_e(b.get("label", ""))}</figcaption>')
                pieces.append('</figure>')
        pieces.append('</body></html>')
        HTML(string=''.join(pieces), base_url=str(mf.parent)).write_pdf(str(out))
        print(f"Draft rendered: {out} | blocks: {len(d['blocks'])} (NOT semantically verified)")
    return 0



def _docx_math_runs(paragraph, parts: list[dict]) -> None:
    """Apply true Word baseline/subscript/superscript properties, never fake underscores."""
    for part in parts:
        key, value = next(iter(part.items()))
        run = paragraph.add_run(str(value))
        if key == "sub":
            run.font.subscript = True
        elif key == "sup":
            run.font.superscript = True


def _build_docx(args: argparse.Namespace) -> int:
    """Create an editable Word reading edition from the SAME reviewed manifest as PDF."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    mf = Path(args.manifest).resolve()
    d = _read_manifest(mf)
    with fitz.open(str(_source_path(d, mf))) as source:
        errors, warnings = _check_manifest(d, source)
        if errors:
            for msg in errors:
                print('ERROR:', msg, file=sys.stderr)
            return 2
        for msg in warnings:
            print('WARNING:', msg, file=sys.stderr)
        out = Path(args.out).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        asset_dir = out.parent / (out.stem + '_draft_assets')
        doc = Document()
        section = doc.sections[0]
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        available_width_cm = 17.0
        styles = doc.styles
        normal = styles['Normal']
        normal.font.name = 'Noto Sans CJK SC'
        normal.font.size = Pt(10.6)
        normal.paragraph_format.line_spacing = 1.48
        normal.paragraph_format.space_after = Pt(8)
        named_styles = []
        for name, size in [('P2CH Title', 17.2), ('P2CH English Title', 10.8), ('P2CH Heading 1', 12.7), ('P2CH Heading 2', 11.5)]:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
            st.font.name = 'Noto Sans CJK SC'
            st.font.color.rgb = RGBColor(22, 23, 24)
            st.font.bold = name != 'P2CH English Title'
            st.font.size = Pt(size)
            st.paragraph_format.space_before = Pt(0 if name in ('P2CH Title', 'P2CH English Title') else 12)
            st.paragraph_format.space_after = Pt(7)
            named_styles.append(st)
        # Explicit Word CJK font fallback. Font files are NEVER packaged in repository.
        for st in [normal] + named_styles:
            rpr = st.element.get_or_add_rPr()
            rfonts = rpr.rFonts
            if rfonts is None:
                rfonts = OxmlElement('w:rFonts')
                rpr.insert(0, rfonts)
            for attr in ('ascii', 'hAnsi', 'eastAsia'):
                rfonts.set(qn('w:' + attr), 'Noto Sans CJK SC')
            spacing = rpr.find(qn('w:spacing'))
            if spacing is None:
                spacing = OxmlElement('w:spacing')
                rpr.append(spacing)
            spacing.set(qn('w:val'), '0')
        meta = d['metadata']
        doc.add_paragraph(str(meta['title_zh']), style='P2CH Title').alignment = WD_ALIGN_PARAGRAPH.LEFT
        doc.add_paragraph(str(meta['title_en']), style='P2CH English Title').alignment = WD_ALIGN_PARAGRAPH.LEFT
        subtitle = ' · '.join(str(meta.get(x)) for x in ('authors','journal','year','doi') if meta.get(x))
        if subtitle:
            p0 = doc.add_paragraph(subtitle)
            p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p0.style = normal
        for i, block in enumerate(d['blocks'], start=1):
            kind = block['type']
            if kind == 'heading':
                p = doc.add_paragraph(str(block['text']), style='P2CH Heading 1' if int(block.get('level',2)) == 2 else 'P2CH Heading 2')
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_with_next = True
            elif kind == 'paragraph':
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.widow_control = True
                if 'runs' in block:
                    for run in block['runs']:
                        key, val = next(iter(run.items()))
                        if key == 'text': p.add_run(str(val))
                        elif key == 'math': _docx_math_runs(p, val)
                else:
                    p.add_run(str(block['text']))
            elif kind == 'note':
                p = doc.add_paragraph('译者核对注：' + str(block['text']))
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif kind == 'references':
                for item in block['items']:
                    p = doc.add_paragraph(str(item))
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(4)
            else:
                img = _crop(source, block, asset_dir, i)
                width_pdf = float(block['bbox'][2]) - float(block['bbox'][0])
                height_pdf = float(block['bbox'][3]) - float(block['bbox'][1])
                # Keep the original aspect ratio; reduce oversized originals to page bounds.
                width_cm = min(available_width_cm, max(2.4, width_pdf * 0.03528))
                if height_pdf > 0:
                    width_cm = min(width_cm, 23.5 * width_pdf / height_pdf)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = bool(block.get('caption') or (kind == 'equation' and block.get('show_label', False)))
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(2)
                p.add_run().add_picture(str(img), width=Cm(width_cm))
                caption = str(block.get('caption', ''))
                if kind != 'equation' and (caption or block.get('label')):
                    cp = doc.add_paragraph((str(block.get('label','')) + ' ' + caption).strip())
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.paragraph_format.space_after = Pt(9)
                elif kind == 'equation' and block.get('show_label', False):
                    cp = doc.add_paragraph(str(block.get('label', '')))
                    cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.core_properties.title = str(meta['title_zh'])
        doc.core_properties.subject = 'Chinese scholarly translation draft; review source alignment before distribution'
        doc.save(str(out))
        print(f'Word draft rendered: {out} | blocks: {len(d["blocks"])} (NOT semantically verified)')
    return 0


def build(args: argparse.Namespace) -> int:
    requested = getattr(args, 'format', None)
    out = Path(args.out).resolve()
    fmt = requested or ('docx' if out.suffix.lower() == '.docx' else 'pdf')
    if fmt not in ('pdf', 'docx', 'both'):
        raise ValueError(f'Unsupported --format: {fmt}')
    if fmt == 'pdf':
        return _build_pdf(argparse.Namespace(manifest=args.manifest, out=str(out)))
    if fmt == 'docx':
        return _build_docx(argparse.Namespace(manifest=args.manifest, out=str(out)))
    stem = out.with_suffix('') if out.suffix.lower() in ('.pdf', '.docx') else out
    for renderer, extension in ((_build_pdf,'.pdf'), (_build_docx,'.docx')):
        code = renderer(argparse.Namespace(manifest=args.manifest, out=str(stem) + extension))
        if code != 0:
            return code
    print('Both outputs generated from one reviewed manifest; review each format independently.')
    return 0

def audit(args: argparse.Namespace) -> int:
    mf = Path(args.manifest).resolve()
    d = _read_manifest(mf)
    errors = []
    warnings = []
    with fitz.open(str(_source_path(d, mf))) as source:
        errs, warns = _check_manifest(d, source)
        errors.extend(errs)
        warnings.extend(warns)
    target = Path(args.pdf)
    if not target.is_file():
        errors.append(f"Built PDF not found: {target}")
    elif target.suffix.lower() == '.docx':
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        wd = Document(str(target))
        text_all = '\n'.join(p.text for p in wd.paragraphs)
        title = str(d.get('metadata', {}).get('title_zh', ''))
        original_title = str(d.get('metadata', {}).get('title_en',''))
        if title not in text_all: errors.append('Chinese title missing from Word text layer')
        if original_title not in text_all: errors.append('Original English title missing from Word text layer')
        if '\ufffd' in text_all: errors.append('Unicode replacement character in Word text')
        expected_images = sum(b.get('type') in {'figure','table','equation'} for b in d.get('blocks',[]))
        if len(wd.inline_shapes) != expected_images:
            errors.append(f'Word images: {len(wd.inline_shapes)}; expected {expected_images}')
        for index, para in enumerate(wd.paragraphs, 1):
            if para.alignment in (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.DISTRIBUTE):
                errors.append(f'Word paragraph {index}: expanded justified/distributed spacing')
        if not any('摘' in para.text for para in wd.paragraphs):
            warnings.append('Cannot locate abstract heading in Word text; check source structure')
        for idx, b in enumerate(d.get('blocks', []), 1):
            if b.get('type') == 'paragraph' and 'runs' in b:
                # The document is structurally checked; scientific semantics still require full visual review.
                expected_tokens = sum('math' in r for r in b['runs'])
                if expected_tokens and not any(r.font.subscript or r.font.superscript for p in wd.paragraphs for r in p.runs):
                    errors.append(f'Word math runs from block {idx} lack actual sub/sup formatting')
                    break
    elif target.suffix.lower() != '.pdf':
        errors.append('Supported audits are PDF and DOCX only')
    else:
        with fitz.open(str(target)) as target_pdf:
            if len(target_pdf) == 0:
                errors.append("Output has no pages")
            total_text = ''
            for i, page in enumerate(target_pdf, start=1):
                content = page.get_text()
                if len(content.strip()) < 6:
                    warnings.append(f"Output page {i}: little extractable text; inspect visually")
                total_text += content
            if str(d.get("metadata", {}).get("title_zh", "")) not in total_text:
                warnings.append("Chinese title not reliably recoverable by text extraction (check CJK font output visually)")
            if '\ufffd' in total_text:
                errors.append("Unicode replacement character found in output extraction")
            if '□' in total_text:
                warnings.append("Square glyph detected; inspect formula/fonts")
    quality = d.get("quality", {})
    if not quality.get("translation_reviewed"):
        warnings.append("Manual full-text translation review NOT marked complete")
    if not quality.get("layout_reviewed"):
        warnings.append("Manual all-page visual review NOT marked complete")
    if getattr(args, "strict", False):
        for i, b in enumerate(d.get("blocks", []), start=1):
            if b.get("type") == "paragraph" and (not b.get("source_page") or not b.get("source_paragraph_id")):
                errors.append(f"Block {i}: strict requires source_page and source_paragraph_id for every paragraph")
        for flag in ("translation_reviewed", "layout_reviewed", "alignment_reviewed", "math_reviewed"):
            if not quality.get(flag):
                errors.append(f"Strict QA requires quality.{flag}=true, set only after actual manual verification")
    report = {"passed_automated": not errors, "errors": errors, "warnings": warnings,
              "semantic_accuracy_certified": False, "manual_review_required": True,
              "strict_requested": bool(getattr(args, "strict", False))}
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not errors else 2


def align(args: argparse.Namespace) -> int:
    """Export a private bilingual traceability sheet. This does not verify semantic fidelity."""
    mf = Path(args.manifest).resolve()
    d = _read_manifest(mf)
    src = _source_path(d, mf)
    out = Path(args.out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    with fitz.open(str(src)) as original, out.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["block", "source_page", "source_paragraph_id", "source_quote", "zh_translation", "check_status"])
        for idx, block in enumerate(d.get("blocks", []), 1):
            if block.get("type") != "paragraph": continue
            page = block.get("source_page", "")
            quote = block.get("source_quote", "")
            state = "needs-manual-check"
            if isinstance(page, int) and 1 <= page <= len(original) and quote:
                source_page_text = re.sub(r"\s+", "", original[page - 1].get_text()).casefold()
                anchor = re.sub(r"\s+", "", str(quote)).casefold()
                state = "anchor-found;not-translation-verified" if anchor in source_page_text else "anchor-not-found;manual-check"
            writer.writerow([idx, page, block.get("source_paragraph_id", ""), quote, _plain_runs(block), state])
    print(f"Private bilingual alignment sheet: {out} (do not publish copyrighted text)")
    return 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(prog="paper2zh", description="PDF/DOCX layout and QA helpers (not an automatic translator)")
    sub = parser.add_subparsers(dest="command", required=True)
    pi = sub.add_parser("inspect", help="Extract blocks and optional source page previews")
    pi.add_argument("pdf")
    pi.add_argument("--out", required=True)
    pi.add_argument("--preview", default=None)
    pi.set_defaults(func=inspect_pdf)
    pb = sub.add_parser("build", help="Render draft PDF, editable DOCX or both from reviewed manuscript JSON")
    pb.add_argument("manifest")
    pb.add_argument("--out", required=True)
    pb.add_argument("--format", choices=["pdf", "docx", "both"], default=None, help="Default inferred from --out suffix; otherwise PDF")
    pb.set_defaults(func=build)
    pa = sub.add_parser("audit", help="Audit PDF or DOCX structure, scientific symbol representation and manual-QA flags")
    pa.add_argument("manifest")
    pa.add_argument("pdf", help="PDF or DOCX output file")
    pa.add_argument("--strict", action="store_true", help="Require manual-reviewed QA flags and paragraph provenance")
    pa.set_defaults(func=audit)
    pl = sub.add_parser("align", help="Export private source-quote-to-translation checking sheet")
    pl.add_argument("manifest")
    pl.add_argument("--out", required=True)
    pl.set_defaults(func=align)
    args = parser.parse_args(argv)
    try:
        return args.func(args)
    except (ValueError, OSError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
